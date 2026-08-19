"""pptx_export CLI — python-pptx **미설치** 경로 (exit 2) + slide-size 파싱.

`tests/test_pptx_export.py` 는 module-level `pytest.importorskip("pptx")` 라
설치 환경(CI)에서 exit-2 경로가 절대 재현되지 않는다(#71 은 `_import_pptx` 를
monkeypatch 해 우회). 여기서는 ImportError 를 내는 `pptx` 스텁을 `PYTHONPATH` 로
주입해, python-pptx 설치 여부와 무관하게 미설치 CLI 동작을 스크립트 직접 실행으로
검증한다. `pptx_export.py` 는 pptx 를 lazy import 하므로 모듈 로드에는 pptx 가 없어도 된다.
"""

import importlib.util
import os
import subprocess
import sys
from pathlib import Path

import pytest

SCRIPT = Path(__file__).parent.parent / "scripts" / "pptx_export.py"

_spec = importlib.util.spec_from_file_location("pptx_export_cli", SCRIPT)
_mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mod)


def test_missing_pptx_cli_exits_two(tmp_path):
    """python-pptx 부재 → exit 2(= partial 계약) + 안내, 트레이스백 없음.

    스텁 `pptx/__init__.py` 가 import 시 ImportError 를 던지고, PYTHONPATH 로
    site-packages 보다 앞에 놓여 실제 설치본을 가린다. 의존성 검사(exit 2)는
    파일 존재/검증(exit 1)보다 먼저라 데이터 파일이 없어도 exit 2 가 나온다.
    """
    stub = tmp_path / "stub" / "pptx"
    stub.mkdir(parents=True)
    (stub / "__init__.py").write_text(
        'raise ImportError("stubbed: python-pptx not installed")\n', encoding="utf-8"
    )
    env = dict(os.environ, PYTHONPATH=str(tmp_path / "stub"))
    result = subprocess.run(
        [sys.executable, str(SCRIPT), str(tmp_path / "any.json")],
        capture_output=True,
        text=True,
        env=env,
    )
    assert result.returncode == 2
    assert "python-pptx" in result.stderr
    assert "Traceback" not in result.stderr


@pytest.mark.parametrize("bad", ["1x", "wide x tall", "1024x"])
def test_parse_slide_size_rejects_malformed(bad):
    with pytest.raises(SystemExit):
        _mod._parse_slide_size(bad)


def test_parse_slide_size_accepts_presets_and_dimensions():
    assert _mod._parse_slide_size("wide") == (1920, 1080)  # SLIDE_PRESETS 조회 분기
    assert _mod._parse_slide_size("auto") is None
    assert _mod._parse_slide_size("1024x768") == (1024, 768)


EXAMPLES = Path(__file__).parent.parent / "examples"


def _cli(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)],
                          capture_output=True, text=True)


def test_multi_input_requires_out():
    """다중 입력은 합칠 덱 경로가 있어야 한다 — 기본 경로 추론은 어느 입력 기준인지 모호하다.

    의존성 검사(exit 2)가 인자 검증(exit 1)보다 먼저라(#71) python-pptx 없는 환경에선
    이 가드에 도달하지 못한다 — 그 순서 자체는 test_missing_pptx_cli_exits_two 가 지킨다.
    """
    pytest.importorskip("pptx")
    r = _cli(EXAMPLES / "microservice-component.json", EXAMPLES / "three-tier-topology.json")
    assert r.returncode == 1
    assert "-o/--out" in r.stderr
    assert "Traceback" not in r.stderr


def test_multi_input_rejects_auto_slide_size(tmp_path):
    """auto 는 덱마다 캔버스가 달라 한 파일에 섞으면 슬라이드 크기가 제각각이 된다."""
    pytest.importorskip("pptx")
    r = _cli(EXAMPLES / "microservice-component.json", EXAMPLES / "three-tier-topology.json",
             "--slide-size", "auto", "-o", tmp_path / "x.pptx")
    assert r.returncode == 1
    assert "auto" in r.stderr
    assert not (tmp_path / "x.pptx").exists()


def test_multi_input_merges_across_views(tmp_path):
    """component + topology + sequence → 한 덱. 슬라이드 수는 각 덱의 합."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    out = tmp_path / "merged.pptx"
    r = _cli(EXAMPLES / "microservice-component.json", EXAMPLES / "three-tier-topology.json",
             EXAMPLES / "order-service-sequence.json", "-o", out)
    assert r.returncode == 0, r.stderr
    assert "병합" in r.stdout
    merged = len(Presentation(str(out)).slides._sldIdLst)

    singles = 0
    for name in ("microservice-component", "three-tier-topology", "order-service-sequence"):
        one = tmp_path / f"{name}.pptx"
        assert _cli(EXAMPLES / f"{name}.json", "-o", one).returncode == 0
        singles += len(Presentation(str(one)).slides._sldIdLst)
    assert merged == singles


def test_multi_input_shares_one_canvas(tmp_path):
    """병합 덱은 캔버스가 하나 — 뷰가 섞여도 슬라이드 크기가 갈리지 않는다."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    out = tmp_path / "merged.pptx"
    assert _cli(EXAMPLES / "microservice-component.json",
                EXAMPLES / "three-tier-topology.json", "-o", out).returncode == 0
    prs = Presentation(str(out))
    assert (prs.slide_width, prs.slide_height) == (18288000, 10287000)   # wide 1920x1080


def test_single_input_unchanged(tmp_path):
    """단일 입력은 -o 없이도 입력 옆에 저장 — 기존 동작 회귀."""
    pytest.importorskip("pptx")
    src = tmp_path / "three-tier-topology.json"
    src.write_text((EXAMPLES / "three-tier-topology.json").read_text(encoding="utf-8"),
                   encoding="utf-8")
    r = _cli(src)
    assert r.returncode == 0, r.stderr
    assert "병합" not in r.stdout
    assert (tmp_path / "three-tier-topology.pptx").exists()
